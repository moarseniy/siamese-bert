from __future__ import annotations

import io
import re
import zipfile
from copy import copy
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


class ExcelInputError(ValueError):
    pass


@dataclass
class ExcelDocument:
    workbook: Workbook
    worksheet: Worksheet
    answers: list[object]
    output_column_index: int


def validate_xlsx(data: bytes, max_uncompressed_bytes: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total_size = sum(item.file_size for item in archive.infolist())
            if total_size > max_uncompressed_bytes:
                raise ExcelInputError(
                    "The uncompressed XLSX content exceeds the configured limit."
                )
            if "[Content_Types].xml" not in archive.namelist():
                raise ExcelInputError("The uploaded file is not a valid XLSX workbook.")
    except zipfile.BadZipFile as exc:
        raise ExcelInputError("The uploaded file is not a valid XLSX workbook.") from exc


def load_excel(
    data: bytes,
    text_column: str,
    output_column: str,
    sheet_name: str | None,
    max_rows: int,
    max_uncompressed_bytes: int,
) -> ExcelDocument:
    validate_xlsx(data, max_uncompressed_bytes)
    try:
        workbook = load_workbook(io.BytesIO(data))
    except Exception as exc:
        raise ExcelInputError(f"Cannot read XLSX workbook: {exc}") from exc
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ExcelInputError(
                f"Sheet {sheet_name!r} not found. Existing: {workbook.sheetnames}"
            )
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.active
    if sheet.max_row > max_rows + 1:
        raise ExcelInputError(
            f"Workbook has {sheet.max_row - 1} data rows; limit is {max_rows}."
        )

    headers: dict[str, list[int]] = {}
    for index, cell in enumerate(sheet[1], start=1):
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        headers.setdefault(name, []).append(index)
    if text_column not in headers:
        raise ExcelInputError(
            f"Column {text_column!r} not found. Existing: {list(headers)}"
        )
    if len(headers[text_column]) != 1:
        raise ExcelInputError(f"Column {text_column!r} occurs more than once.")

    if output_column in headers:
        if len(headers[output_column]) != 1:
            raise ExcelInputError(f"Column {output_column!r} occurs more than once.")
        output_index = headers[output_column][0]
    else:
        output_index = sheet.max_column + 1
        output_cell = sheet.cell(1, output_index, output_column)
        if output_index > 1:
            source_cell = sheet.cell(1, output_index - 1)
            output_cell._style = copy(source_cell._style)
            output_cell.font = copy(source_cell.font)
            output_cell.fill = copy(source_cell.fill)
            output_cell.border = copy(source_cell.border)
            output_cell.alignment = copy(source_cell.alignment)
            output_cell.number_format = source_cell.number_format
            output_cell.protection = copy(source_cell.protection)
            source_letter = source_cell.column_letter
            output_letter = output_cell.column_letter
            sheet.column_dimensions[output_letter].width = max(
                18, sheet.column_dimensions[source_letter].width or 0
            )

    text_index = headers[text_column][0]
    answers = [
        sheet.cell(row_number, text_index).value
        for row_number in range(2, sheet.max_row + 1)
    ]
    return ExcelDocument(workbook, sheet, answers, output_index)


def write_predictions(document: ExcelDocument, predictions: list[str]) -> bytes:
    if len(document.answers) != len(predictions):
        raise ValueError("Prediction count does not match the workbook row count.")
    for row_number, prediction in enumerate(predictions, start=2):
        document.worksheet.cell(
            row_number, document.output_column_index, prediction
        )
    output = io.BytesIO()
    document.workbook.save(output)
    document.workbook.close()
    return output.getvalue()


def build_result_zip(
    result_xlsx: bytes,
    result_name: str,
    codebook_bytes: bytes,
) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(result_name, result_xlsx)
        archive.writestr("codebook.xlsx", codebook_bytes)
    return output.getvalue()


def safe_result_name(upload_name: str | None) -> str:
    raw_name = (upload_name or "responses.xlsx").replace("\\", "/")
    name = Path(raw_name.rsplit("/", 1)[-1]).name
    stem = re.sub(r"[^\w.-]+", "_", Path(name).stem, flags=re.UNICODE).strip("._")
    stem = stem or "responses"
    return f"{stem}_predictions.xlsx"
