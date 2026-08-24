from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


LEGACY_FORMAT = "category_subcategory_v1"
TEXT_ONLY_FORMAT = "category_subcategory_v2"
SUPPORTED_FORMATS = {LEGACY_FORMAT, TEXT_ONLY_FORMAT}

_CYRILLIC_TO_LATIN = str.maketrans(
    {
        "А": "A",
        "В": "B",
        "С": "C",
        "Е": "E",
        "К": "K",
        "М": "M",
        "Н": "H",
        "О": "O",
        "Р": "P",
        "Т": "T",
        "Х": "X",
    }
)


def clean_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalize_code(value: object) -> str:
    code = clean_text(value).upper().translate(_CYRILLIC_TO_LATIN)
    return re.sub(r"\s+", "", code)


def parent_code(code: str) -> str:
    match = re.match(r"^([A-Z]+)", code)
    return match.group(1) if match else code


@dataclass(frozen=True)
class CodeEntry:
    code: str
    category: str
    subcategory: str
    description: str


def load_codebook(path: Path, description_format: str) -> list[CodeEntry]:
    if description_format not in SUPPORTED_FORMATS:
        raise ValueError(
            f"Unsupported code_description_format: {description_format!r}."
        )
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise FileNotFoundError(f"codebook.xlsx not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers: dict[str, int] = {}
    for index, cell in enumerate(sheet[1], start=1):
        name = clean_text(cell.value)
        if name:
            headers[name] = index
    required = ["Код", "Категория", "Подкатегория"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(
            f"Missing codebook columns: {missing}. Existing: {list(headers)}"
        )

    entries: list[CodeEntry] = []
    seen: set[str] = set()
    parent_names: dict[str, str] = {}
    for row_number in range(2, sheet.max_row + 1):
        code = normalize_code(sheet.cell(row_number, headers["Код"]).value)
        category = clean_text(sheet.cell(row_number, headers["Категория"]).value)
        subcategory = clean_text(
            sheet.cell(row_number, headers["Подкатегория"]).value
        )
        if not code and not category and not subcategory:
            continue
        if not code or not category or not subcategory:
            raise ValueError(
                f"Invalid codebook row {row_number}: all three fields are required."
            )
        parent = parent_code(code)
        if parent == code:
            raise ValueError(f"Invalid leaf code {code!r} at row {row_number}.")
        if code in seen:
            raise ValueError(f"Duplicate code in codebook: {code}")
        if parent in parent_names and parent_names[parent] != category:
            raise ValueError(f"Conflicting categories for parent code {parent}.")
        seen.add(code)
        parent_names[parent] = category
        description = f"Категория: {category}. Подкатегория: {subcategory}"
        if description_format == LEGACY_FORMAT:
            description = f"{code}. {description}"
        entries.append(
            CodeEntry(
                code=code,
                category=category,
                subcategory=subcategory,
                description=description,
            )
        )
    workbook.close()
    if not entries:
        raise ValueError("Codebook contains no leaf codes.")
    return entries

