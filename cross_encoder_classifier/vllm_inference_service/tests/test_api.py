import asyncio
import io
import zipfile
from types import SimpleNamespace

import httpx
from openpyxl import Workbook, load_workbook

from app.main import app


class FakePredictor:
    async def predict(self, answers):
        assert answers == ["Низкая зарплата"]
        return ["A1:2"]


class FakeSemaphore:
    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, traceback):
        return False


def make_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Ответ"])
    sheet.append(["Низкая зарплата"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_predict_endpoint_returns_zip() -> None:
    app.state.settings = SimpleNamespace(
        max_upload_bytes=10_000_000,
        text_column="Ответ",
        output_column="Предсказание",
        sheet_name=None,
        max_rows=100,
        max_uncompressed_xlsx_bytes=10_000_000,
    )
    app.state.predictor = FakePredictor()
    app.state.codebook_bytes = b"codebook-content"
    app.state.request_semaphore = FakeSemaphore()

    async def request() -> httpx.Response:
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(
            transport=transport, base_url="http://test"
        ) as client:
            return await client.post(
                "/predict",
                files={
                    "file": (
                        "answers.xlsx",
                        make_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )

    response = asyncio.run(request())

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.read("codebook.xlsx") == b"codebook-content"
        result = load_workbook(
            io.BytesIO(archive.read("answers_predictions.xlsx"))
        ).active
    assert result.cell(1, 2).value == "Предсказание"
    assert result.cell(2, 2).value == "A1:2"
