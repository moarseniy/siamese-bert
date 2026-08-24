import io
import zipfile

from openpyxl import Workbook, load_workbook

from app.excel_io import (
    build_result_zip,
    load_excel,
    safe_result_name,
    write_predictions,
)


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Опрос"
    sheet.append(["id", "Ответ"])
    sheet.append([1, "Низкая зарплата"])
    sheet.append([2, "Хороший коллектив"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_adds_prediction_column_and_zip() -> None:
    document = load_excel(
        workbook_bytes(),
        text_column="Ответ",
        output_column="Предсказание",
        sheet_name=None,
        max_rows=100,
        max_uncompressed_bytes=10_000_000,
    )
    result = write_predictions(document, ["A1:2", "B1:1"])
    archive_bytes = build_result_zip(result, "answers_predictions.xlsx", b"codebook")

    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        assert set(archive.namelist()) == {
            "answers_predictions.xlsx",
            "codebook.xlsx",
        }
        result_workbook = load_workbook(
            io.BytesIO(archive.read("answers_predictions.xlsx"))
        )
    sheet = result_workbook["Опрос"]
    assert sheet.cell(1, 3).value == "Предсказание"
    assert sheet.cell(2, 3).value == "A1:2"
    assert sheet.cell(3, 3).value == "B1:1"


def test_result_name_cannot_escape_zip_directory() -> None:
    assert safe_result_name(r"..\folder\ответ.xlsx") == "ответ_predictions.xlsx"
