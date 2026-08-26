from __future__ import annotations

import asyncio
import io
import json
import re
import zipfile
from contextlib import asynccontextmanager
from pathlib import Path
from urllib.parse import quote

import httpx
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from tokenizers import Tokenizer
import yaml


SERVICE_CONFIG_PATH = Path(__file__).resolve().parents[1] / "config.yaml"
SETTINGS = yaml.safe_load(SERVICE_CONFIG_PATH.read_text(encoding="utf-8"))
VLLM_URL = str(SETTINGS["vllm_url"]).rstrip("/")
VLLM_API_KEY = str(SETTINGS.get("vllm_api_key") or "").strip()
SERVED_MODEL_NAME = str(SETTINGS.get("served_model_name") or "").strip()
TEXT_COLUMN = str(SETTINGS.get("text_column", "Ответ"))
OUTPUT_COLUMN = str(SETTINGS.get("output_column", "Предсказание"))
SHEET_NAME = str(SETTINGS.get("sheet_name") or "").strip()
PAIR_BATCH_SIZE = int(SETTINGS.get("pair_batch_size", 512))
REQUEST_TIMEOUT = float(SETTINGS.get("request_timeout", 600))
CODEBOOK_PATH = Path(SETTINGS.get("codebook_path", "artifacts/codebook.xlsx"))
TOKENIZER_PATH = Path(SETTINGS.get("tokenizer_path", "artifacts/tokenizer.json"))
CONFIG_PATH = Path(
    SETTINGS.get("classifier_config_path", "artifacts/classifier_config.json")
)


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_assets() -> tuple[list[tuple[str, str]], Tokenizer, dict[str, object]]:
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    description_format = config.get(
        "code_description_format", "category_subcategory_v2"
    )

    workbook = load_workbook(CODEBOOK_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    headers = {text(cell.value): index for index, cell in enumerate(sheet[1], 1)}
    required = {"Код", "Категория", "Подкатегория"}
    if missing := required - headers.keys():
        raise ValueError(f"Missing codebook columns: {sorted(missing)}")

    codebook: list[tuple[str, str]] = []
    for row in range(2, sheet.max_row + 1):
        code = re.sub(r"\s+", "", text(sheet.cell(row, headers["Код"]).value).upper())
        category = text(sheet.cell(row, headers["Категория"]).value)
        subcategory = text(sheet.cell(row, headers["Подкатегория"]).value)
        if not code and not category and not subcategory:
            continue
        if not code or not category or not subcategory:
            raise ValueError(f"Invalid codebook row {row}")
        description = f"Категория: {category}. Подкатегория: {subcategory}"
        if description_format in {
            "category_subcategory_v1",
            "code_category_subcategory_v1",
        }:
            description = f"{code}. {description}"
        codebook.append((code, description))
    workbook.close()
    if not codebook:
        raise ValueError("Codebook is empty")

    tokenizer = Tokenizer.from_file(str(TOKENIZER_PATH))
    tokenizer.no_padding()
    tokenizer.enable_truncation(
        max_length=int(config.get("max_length", 256)), strategy="longest_first"
    )
    return codebook, tokenizer, config


def add_prefix(value: object, prefix: str) -> str:
    value = text(value)
    if not prefix or ";" not in value:
        return value
    before, after = value.split(";", 1)
    return f"{before.strip()}; {prefix} {after.strip()}".strip()


async def classify(app: FastAPI, token_ids: list[list[int]]) -> list[list[float]]:
    response = await app.state.client.post(
        f"{VLLM_URL}/classify",
        json={"model": app.state.model, "input": token_ids, "use_activation": True},
    )
    if not response.is_success:
        raise RuntimeError(f"vLLM HTTP {response.status_code}: {response.text[:500]}")
    try:
        items = sorted(response.json()["data"], key=lambda item: item["index"])
        probabilities = [list(map(float, item["probs"])) for item in items]
    except (KeyError, TypeError, ValueError) as exc:
        raise RuntimeError("Invalid response from vLLM /classify") from exc
    if len(probabilities) != len(token_ids) or any(
        len(values) != 4 for values in probabilities
    ):
        raise RuntimeError("vLLM returned an unexpected number of predictions")
    return probabilities


def decode(
    probabilities: list[list[float]],
    codebook: list[tuple[str, str]],
    threshold: float,
    max_labels: int,
) -> str:
    candidates: list[tuple[float, str, int]] = []
    for values, (code, _) in zip(probabilities, codebook, strict=True):
        presence = 1.0 - values[0]
        if presence >= threshold:
            sentiment = max(range(3), key=lambda index: values[index + 1])
            candidates.append((presence, code, sentiment))
    candidates.sort(reverse=True)
    return ", ".join(
        f"{code}:{sentiment}"
        for _, code, sentiment in candidates[:max_labels]
    ) or "UNKNOWN"


async def predict_answers(app: FastAPI, answers: list[object]) -> list[str]:
    codebook = app.state.codebook
    config = app.state.config
    results = ["UNKNOWN"] * len(answers)
    rows_per_group = max(1, PAIR_BATCH_SIZE // len(codebook))
    prefix = text(config.get("after_semicolon_prefix", ""))

    for start in range(0, len(answers), rows_per_group):
        pairs: list[tuple[str, str]] = []
        active_rows: list[int] = []
        for row, answer in enumerate(answers[start : start + rows_per_group]):
            answer = add_prefix(answer, prefix)
            if answer:
                active_rows.append(row)
                pairs.extend((answer, description) for _, description in codebook)

        probabilities: list[list[float]] = []
        for offset in range(0, len(pairs), PAIR_BATCH_SIZE):
            encodings = app.state.tokenizer.encode_batch(
                pairs[offset : offset + PAIR_BATCH_SIZE], add_special_tokens=True
            )
            probabilities.extend(
                await classify(app, [encoding.ids for encoding in encodings])
            )

        for position, row in enumerate(active_rows):
            left = position * len(codebook)
            results[start + row] = decode(
                probabilities[left : left + len(codebook)],
                codebook,
                float(config["threshold"]),
                int(config.get("max_labels", 6)),
            )
    return results


@asynccontextmanager
async def lifespan(app: FastAPI):
    codebook, tokenizer, config = load_assets()
    headers = {"Authorization": f"Bearer {VLLM_API_KEY}"} if VLLM_API_KEY else {}
    client = httpx.AsyncClient(headers=headers, timeout=REQUEST_TIMEOUT)
    response = await client.get(f"{VLLM_URL}/v1/models")
    response.raise_for_status()
    models = response.json().get("data", [])
    model = SERVED_MODEL_NAME or (models[0].get("id") if models else "")
    if not model:
        raise RuntimeError("vLLM returned no served model")

    app.state.client = client
    app.state.model = model
    app.state.codebook = codebook
    app.state.tokenizer = tokenizer
    app.state.config = config
    app.state.codebook_bytes = CODEBOOK_PATH.read_bytes()
    app.state.lock = asyncio.Lock()
    try:
        yield
    finally:
        await client.aclose()


app = FastAPI(docs_url=None, redoc_url=None, openapi_url=None, lifespan=lifespan)


@app.post("/predict")
async def predict(file: UploadFile = File(...)) -> Response:
    filename = file.filename or "responses.xlsx"
    if Path(filename).suffix.lower() != ".xlsx":
        raise HTTPException(400, "Only .xlsx files are accepted")

    try:
        workbook = load_workbook(io.BytesIO(await file.read()))
        sheet = workbook[SHEET_NAME] if SHEET_NAME else workbook.active
        headers = {text(cell.value): index for index, cell in enumerate(sheet[1], 1)}
        if TEXT_COLUMN not in headers:
            raise ValueError(f"Column {TEXT_COLUMN!r} not found")
        output_column = headers.get(OUTPUT_COLUMN, sheet.max_column + 1)
        sheet.cell(1, output_column, OUTPUT_COLUMN)
        answers = [
            sheet.cell(row, headers[TEXT_COLUMN]).value
            for row in range(2, sheet.max_row + 1)
        ]
        async with app.state.lock:
            predictions = await predict_answers(app, answers)
        for row, prediction in enumerate(predictions, 2):
            sheet.cell(row, output_column, prediction)

        result = io.BytesIO()
        workbook.save(result)
        workbook.close()
        stem = re.sub(r"[^\w.-]+", "_", Path(filename).stem).strip("._") or "result"
        result_name = f"{stem}_predictions.xlsx"
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as output:
            output.writestr(result_name, result.getvalue())
            output.writestr("codebook.xlsx", app.state.codebook_bytes)
    except (ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise HTTPException(400, str(exc)) from exc
    except (httpx.HTTPError, RuntimeError) as exc:
        raise HTTPException(502, str(exc)) from exc

    name = quote(f"{Path(result_name).stem}.zip")
    return Response(
        archive.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{name}"},
    )
